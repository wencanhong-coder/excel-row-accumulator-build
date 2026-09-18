# -*- coding: utf-8 -*-
"""
Excel 行累加工具

作用：
1. 按指定的“分组列”判断哪些行属于同一条记录；
2. 对指定的“累加列”进行数值求和；
3. 其余字段可保留首个非空值，或将不同文本去重后拼接；
4. 输出一个新的 Excel 文件，包含“汇总结果”“汇总说明”“字段冲突”三个工作表。

直接双击或运行：
    python Excel行累加.py

命令行示例：
    python Excel行累加.py --input "原始数据.xlsx" --output "汇总结果.xlsx" \
        --sheet "明细" --header-row 1 --group-by "项目,规格" \
        --sum-cols "数量,金额" --text-policy join

依赖：
    pip install openpyxl

说明：
- 仅支持 .xlsx 文件；
- 为避免误改原始数据，脚本不会覆盖输入文件；
- 求和列遇到公式或非数值文本会停止，并提示具体位置；
- 空白值不等于 0：若一组记录在某个求和列中全部为空，汇总结果仍为空。
"""

from __future__ import annotations

import argparse
import copy
import sys
from collections import OrderedDict
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Dict, Iterable, List, Sequence, Tuple

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ModuleNotFoundError:
    print("未检测到 openpyxl。请先在命令提示符中执行：")
    print("pip install openpyxl")
    sys.exit(1)


class AggregationError(Exception):
    """用于向用户输出清晰、可操作的错误信息。"""


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def display_value(value: Any) -> str:
    """将单元格值转换为适合审核表显示的文本。"""
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    return str(value).strip() if isinstance(value, str) else str(value)


def value_token(value: Any) -> Tuple[str, str]:
    """用于去重和分组；文本前后空格不影响判断。"""
    if is_blank(value):
        return ("blank", "")
    if isinstance(value, (datetime, date, time)):
        return (type(value).__name__, display_value(value))
    if isinstance(value, str):
        return ("text", value.strip())
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return ("number", str(Decimal(str(value)).normalize()))
    return (type(value).__name__, str(value))


def parse_decimal(value: Any, sheet_name: str, row_number: int, header: str) -> Decimal | None:
    """把数值单元格转成 Decimal；空白返回 None，以保留“全空”语义。"""
    if is_blank(value):
        return None
    if isinstance(value, bool):
        raise AggregationError(
            f"工作表“{sheet_name}”第 {row_number} 行“{header}”列是布尔值，不能作为求和数值。"
        )
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(value))
        except InvalidOperation as exc:
            raise AggregationError(
                f"工作表“{sheet_name}”第 {row_number} 行“{header}”列的数值无法读取：{value!r}"
            ) from exc
    if isinstance(value, str):
        cleaned = value.strip().replace(",", "").replace("，", "")
        try:
            return Decimal(cleaned)
        except InvalidOperation as exc:
            raise AggregationError(
                f"工作表“{sheet_name}”第 {row_number} 行“{header}”列为“{value}”，"
                "不是可累加的数值。请先将该列清洗为纯数字。"
            ) from exc
    raise AggregationError(
        f"工作表“{sheet_name}”第 {row_number} 行“{header}”列的类型不支持求和：{type(value).__name__}。"
    )


def decimal_to_excel(value: Decimal) -> int | float:
    """尽量保持整数为整数，避免输出 5.0 这类不必要的小数。"""
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def split_column_names(raw: str) -> List[str]:
    """支持中文或英文逗号。"""
    return [item.strip() for item in raw.replace("，", ",").split(",") if item.strip()]


def clean_path_text(raw: str) -> Path:
    """允许用户从资源管理器复制带引号的文件路径。"""
    return Path(raw.strip().strip('"').strip("'")).expanduser()


def validate_excel_path(path: Path, label: str, must_exist: bool) -> None:
    if must_exist and not path.is_file():
        raise AggregationError(f"{label}不存在：{path}")
    if path.suffix.lower() != ".xlsx":
        raise AggregationError(f"{label}必须是 .xlsx 文件：{path.name}")


def find_last_data_row(ws, header_row: int, last_column: int) -> int:
    for row_number in range(ws.max_row, header_row, -1):
        if any(not is_blank(ws.cell(row_number, column).value) for column in range(1, last_column + 1)):
            return row_number
    return header_row


def get_headers(ws, header_row: int) -> Tuple[List[str], Dict[str, int]]:
    headers: List[str] = []
    header_map: Dict[str, int] = {}
    nonempty_columns = [
        column
        for column in range(1, ws.max_column + 1)
        if not is_blank(ws.cell(header_row, column).value)
    ]
    if not nonempty_columns:
        raise AggregationError(f"工作表“{ws.title}”第 {header_row} 行没有可识别的表头。")
    last_header_column = max(nonempty_columns)
    for column in range(1, last_header_column + 1):
        value = ws.cell(header_row, column).value
        if is_blank(value):
            raise AggregationError(
                f"工作表“{ws.title}”第 {header_row} 行第 {column} 列的表头为空。"
                "请删除该空列，或补齐唯一的列名后再运行。"
            )
        name = str(value).strip()
        if name in header_map:
            raise AggregationError(f"表头第 {header_row} 行存在重复列名“{name}”，请先修改为唯一名称。")
        headers.append(name)
        header_map[name] = column - 1
    return headers, header_map


def validate_requested_columns(
    names: Sequence[str], header_map: Dict[str, int], usage: str
) -> List[int]:
    missing = [name for name in names if name not in header_map]
    if missing:
        raise AggregationError(
            f"{usage}中找不到列：{'、'.join(missing)}。可用列：{'、'.join(header_map)}"
        )
    if len(set(names)) != len(names):
        raise AggregationError(f"{usage}中存在重复列名，请删除重复项。")
    return [header_map[name] for name in names]


def ensure_no_formulas(ws, first_row: int, last_row: int, column_count: int) -> None:
    for row_number in range(first_row, last_row + 1):
        for column in range(1, column_count + 1):
            value = ws.cell(row_number, column).value
            if isinstance(value, str) and value.startswith("="):
                header = ws.cell(first_row - 1, column).value or get_column_letter(column)
                raise AggregationError(
                    f"工作表“{ws.title}”第 {row_number} 行“{header}”列包含公式。"
                    "请先在 Excel 中复制该表，并选择“选择性粘贴→数值”，再运行本脚本。"
                )


def copy_cell_style(source, target) -> None:
    """复制结果单元格的基本显示格式，不复制公式和数据验证。"""
    if source.has_style:
        target._style = copy.copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.alignment:
        target.alignment = copy.copy(source.alignment)
    if source.font:
        target.font = copy.copy(source.font)
    if source.fill:
        target.fill = copy.copy(source.fill)
    if source.border:
        target.border = copy.copy(source.border)
    if source.protection:
        target.protection = copy.copy(source.protection)


def make_unique_sheet_name(workbook: Workbook, preferred: str) -> str:
    if preferred not in workbook.sheetnames:
        return preferred
    sequence = 2
    while f"{preferred}{sequence}" in workbook.sheetnames:
        sequence += 1
    return f"{preferred}{sequence}"


def column_width(value: Any) -> float:
    text = display_value(value)
    width = 0.0
    for char in text:
        width += 2.0 if ord(char) > 127 else 1.0
    return min(max(width + 2, 10), 45)


def write_title(ws, title: str, last_column: int) -> None:
    ws.cell(1, 1, title)
    if last_column > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    cell = ws.cell(1, 1)
    cell.font = Font(name="微软雅黑", size=14, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 26


def create_summary_sheet(
    output_wb: Workbook,
    source_ws,
    headers: Sequence[str],
    records: Sequence[Dict[str, Any]],
    source_path: Path,
    header_row: int,
    group_names: Sequence[str],
    sum_names: Sequence[str],
    text_policy: str,
) -> str:
    sheet_name = make_unique_sheet_name(output_wb, "汇总结果")
    ws = output_wb.create_sheet(sheet_name, 0)
    total_columns = len(headers) + 1
    layout_columns = max(total_columns, 8)
    write_title(ws, "Excel 行累加结果", layout_columns)

    ws.cell(2, 1, "源文件")
    ws.cell(2, 2, source_path.name)
    ws.cell(2, 4, "源工作表")
    ws.cell(2, 5, source_ws.title)
    ws.cell(3, 1, "分组列")
    ws.cell(3, 2, "、".join(group_names))
    ws.cell(3, 4, "累加列")
    ws.cell(3, 5, "、".join(sum_names) if sum_names else "无（仅合并行）")
    ws.cell(3, 7, "其他字段")
    ws.cell(3, 8, "去重拼接" if text_policy == "join" else "保留首个非空值")

    header_output_row = 4
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(name="微软雅黑", bold=True, color="17365D")
    for output_column, header in enumerate(headers, start=1):
        source_cell = source_ws.cell(header_row, output_column)
        target_cell = ws.cell(header_output_row, output_column, header)
        copy_cell_style(source_cell, target_cell)
        target_cell.fill = header_fill
        target_cell.font = header_font
        target_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    count_cell = ws.cell(header_output_row, len(headers) + 1, "合并行数")
    count_cell.fill = header_fill
    count_cell.font = header_font
    count_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for output_row, record in enumerate(records, start=header_output_row + 1):
        values = record["values"]
        styles = record["styles"]
        for output_column, value in enumerate(values, start=1):
            target = ws.cell(output_row, output_column, value)
            copy_cell_style(styles[output_column - 1], target)
            target.alignment = copy.copy(styles[output_column - 1].alignment)
        target = ws.cell(output_row, len(headers) + 1, record["row_count"])
        target.alignment = Alignment(horizontal="center", vertical="center")

    data_last_row = header_output_row + len(records)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{header_output_row}:{get_column_letter(total_columns)}{data_last_row}"
    ws.row_dimensions[header_output_row].height = 24

    for column in range(1, total_columns + 1):
        sample_values = [ws.cell(header_output_row, column).value]
        sample_values.extend(ws.cell(row, column).value for row in range(header_output_row + 1, min(data_last_row, header_output_row + 30) + 1))
        ws.column_dimensions[get_column_letter(column)].width = min(
            45, max(column_width(value) for value in sample_values)
        )
    for column in range(total_columns + 1, layout_columns + 1):
        ws.column_dimensions[get_column_letter(column)].width = 16
    return sheet_name


def create_info_sheet(
    output_wb: Workbook,
    source_path: Path,
    source_sheet_name: str,
    header_row: int,
    group_names: Sequence[str],
    sum_names: Sequence[str],
    text_policy: str,
    original_rows: int,
    result_rows: int,
    conflict_count: int,
) -> str:
    sheet_name = make_unique_sheet_name(output_wb, "汇总说明")
    ws = output_wb.create_sheet(sheet_name)
    write_title(ws, "本次汇总说明", 2)
    rows = [
        ("源文件", source_path.name),
        ("源工作表", source_sheet_name),
        ("源表头行", header_row),
        ("分组列", "、".join(group_names)),
        ("累加列", "、".join(sum_names) if sum_names else "无（仅合并行）"),
        ("其他字段处理", "去重拼接" if text_policy == "join" else "保留首个非空值"),
        ("原始非空数据行数", original_rows),
        ("汇总后行数", result_rows),
        ("减少行数", original_rows - result_rows),
        ("字段冲突数", conflict_count),
        ("重要提示", "原始文件未被修改；请复核“字段冲突”工作表中的记录。"),
    ]
    for row_number, (label, value) in enumerate(rows, start=3):
        label_cell = ws.cell(row_number, 1, label)
        label_cell.font = Font(name="微软雅黑", bold=True, color="17365D")
        label_cell.fill = PatternFill("solid", fgColor="D9EAF7")
        label_cell.alignment = Alignment(vertical="top", wrap_text=True)
        value_cell = ws.cell(row_number, 2, value)
        value_cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 70
    ws.freeze_panes = "A3"
    return sheet_name


def create_conflict_sheet(
    output_wb: Workbook,
    conflicts: Sequence[Dict[str, Any]],
) -> str:
    sheet_name = make_unique_sheet_name(output_wb, "字段冲突")
    ws = output_wb.create_sheet(sheet_name)
    write_title(ws, "字段冲突审核", 5)
    headers = ["分组键", "冲突字段", "保留或输出值", "该组所有不同值", "原始行号"]
    header_fill = PatternFill("solid", fgColor="FCE4D6")
    header_font = Font(name="微软雅黑", bold=True, color="9C0006")
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(3, column, header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if conflicts:
        for row_number, conflict in enumerate(conflicts, start=4):
            values = [
                conflict["group_key"],
                conflict["field"],
                conflict["kept_value"],
                conflict["all_values"],
                conflict["source_rows"],
            ]
            for column, value in enumerate(values, start=1):
                cell = ws.cell(row_number, column, value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    else:
        ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=5)
        ws.cell(4, 1, "未发现非分组、非累加字段存在不一致的合并记录。")
        ws.cell(4, 1).alignment = Alignment(vertical="center")

    for column, width in enumerate([38, 20, 30, 55, 18], start=1):
        ws.column_dimensions[get_column_letter(column)].width = width
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:E{max(4, 3 + len(conflicts))}"
    return sheet_name


def aggregate_rows(
    source_ws,
    header_row: int,
    group_names: Sequence[str],
    sum_names: Sequence[str],
    text_policy: str,
) -> Tuple[List[str], List[Dict[str, Any]], List[Dict[str, Any]], int]:
    headers, header_map = get_headers(source_ws, header_row)
    group_indexes = validate_requested_columns(group_names, header_map, "分组列")
    sum_indexes = validate_requested_columns(sum_names, header_map, "累加列")
    overlap = set(group_indexes) & set(sum_indexes)
    if overlap:
        overlap_names = [headers[index] for index in sorted(overlap)]
        raise AggregationError(f"同一列不能同时作为分组列和累加列：{'、'.join(overlap_names)}")

    last_column = len(headers)
    last_data_row = find_last_data_row(source_ws, header_row, last_column)
    if last_data_row <= header_row:
        raise AggregationError(f"工作表“{source_ws.title}”的表头下方没有可汇总的数据。")
    ensure_no_formulas(source_ws, header_row + 1, last_data_row, last_column)

    grouped: "OrderedDict[Tuple[Tuple[str, str], ...], Dict[str, Any]]" = OrderedDict()
    source_row_count = 0

    for row_number in range(header_row + 1, last_data_row + 1):
        cells = [source_ws.cell(row_number, column) for column in range(1, last_column + 1)]
        row_values = [cell.value for cell in cells]
        if all(is_blank(value) for value in row_values):
            continue
        source_row_count += 1
        group_key = tuple(value_token(row_values[index]) for index in group_indexes)

        if group_key not in grouped:
            grouped[group_key] = {
                "values": list(row_values),
                "styles": list(cells),
                "row_count": 0,
                "source_rows": [],
                "sum_totals": {index: Decimal("0") for index in sum_indexes},
                "sum_has_value": {index: False for index in sum_indexes},
                "other_values": {index: [] for index in range(last_column) if index not in group_indexes and index not in sum_indexes},
            }

        record = grouped[group_key]
        record["row_count"] += 1
        record["source_rows"].append(row_number)

        for index in sum_indexes:
            number = parse_decimal(row_values[index], source_ws.title, row_number, headers[index])
            if number is not None:
                record["sum_totals"][index] += number
                record["sum_has_value"][index] = True

        for index in record["other_values"]:
            value = row_values[index]
            if is_blank(value):
                continue
            token = value_token(value)
            if all(existing[0] != token for existing in record["other_values"][index]):
                record["other_values"][index].append((token, value))

    records: List[Dict[str, Any]] = []
    conflicts: List[Dict[str, Any]] = []
    for record in grouped.values():
        values = list(record["values"])
        for index in sum_indexes:
            values[index] = (
                decimal_to_excel(record["sum_totals"][index])
                if record["sum_has_value"][index]
                else None
            )

        group_key_text = "；".join(
            f"{headers[index]}={display_value(values[index])}" for index in group_indexes
        )
        source_rows_text = "、".join(str(number) for number in record["source_rows"])
        for index, distinct_values in record["other_values"].items():
            if not distinct_values:
                values[index] = None
                continue
            raw_values = [item[1] for item in distinct_values]
            if text_policy == "join":
                values[index] = "；".join(display_value(value) for value in raw_values)
            else:
                values[index] = raw_values[0]

            if len(raw_values) > 1:
                conflicts.append(
                    {
                        "group_key": group_key_text,
                        "field": headers[index],
                        "kept_value": display_value(values[index]),
                        "all_values": "；".join(display_value(value) for value in raw_values),
                        "source_rows": source_rows_text,
                    }
                )

        record["values"] = values
        records.append(record)

    return headers, records, conflicts, source_row_count


def run_aggregation(
    input_path: Path,
    output_path: Path,
    sheet_name: str,
    header_row: int,
    group_names: Sequence[str],
    sum_names: Sequence[str],
    text_policy: str,
    overwrite: bool,
) -> Dict[str, Any]:
    validate_excel_path(input_path, "输入文件", must_exist=True)
    validate_excel_path(output_path, "输出文件", must_exist=False)
    if input_path.resolve() == output_path.resolve():
        raise AggregationError("输出文件不能与输入文件相同，以免覆盖原始数据。")
    if output_path.exists() and not overwrite:
        raise AggregationError(
            f"输出文件已存在：{output_path.name}。请换一个文件名，或在命令行中加入 --overwrite。"
        )
    if header_row < 1:
        raise AggregationError("表头行号必须大于等于 1。")
    if not group_names:
        raise AggregationError("请至少填写一个分组列，用于判断哪些行需要合并。")

    input_wb = load_workbook(input_path, data_only=False, read_only=False)
    if sheet_name not in input_wb.sheetnames:
        raise AggregationError(
            f"找不到工作表“{sheet_name}”。可用工作表：{'、'.join(input_wb.sheetnames)}"
        )
    source_ws = input_wb[sheet_name]
    headers, records, conflicts, original_row_count = aggregate_rows(
        source_ws, header_row, group_names, sum_names, text_policy
    )

    output_wb = Workbook()
    output_wb.remove(output_wb.active)
    result_sheet = create_summary_sheet(
        output_wb,
        source_ws,
        headers,
        records,
        input_path,
        header_row,
        group_names,
        sum_names,
        text_policy,
    )
    info_sheet = create_info_sheet(
        output_wb,
        input_path,
        source_ws.title,
        header_row,
        group_names,
        sum_names,
        text_policy,
        original_row_count,
        len(records),
        len(conflicts),
    )
    conflict_sheet = create_conflict_sheet(output_wb, conflicts)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(output_path)
    return {
        "result_sheet": result_sheet,
        "info_sheet": info_sheet,
        "conflict_sheet": conflict_sheet,
        "original_rows": original_row_count,
        "result_rows": len(records),
        "conflicts": len(conflicts),
    }


def prompt_text(prompt: str, default: str | None = None, required: bool = True) -> str:
    suffix = f"（直接回车使用：{default}）" if default else ""
    while True:
        value = input(f"{prompt}{suffix}：").strip()
        if value:
            return value
        if default is not None:
            return default
        if not required:
            return ""
        print("此项不能为空，请重新输入。")


def prompt_header_row(default: int = 1) -> int:
    while True:
        raw = prompt_text("表头所在行号", str(default))
        try:
            value = int(raw)
            if value >= 1:
                return value
        except ValueError:
            pass
        print("请输入大于等于 1 的整数。")


def interactive_config() -> Dict[str, Any]:
    print("\nExcel 行累加工具")
    print("相同分组列的记录会合并；数值列可累加；原文件不会被修改。\n")
    input_path = clean_path_text(prompt_text("请输入输入 Excel 文件完整路径"))
    validate_excel_path(input_path, "输入文件", must_exist=True)

    preview_wb = load_workbook(input_path, data_only=False, read_only=False)
    print(f"可用工作表：{'、'.join(preview_wb.sheetnames)}")
    sheet_name = prompt_text("请输入要汇总的工作表名称", preview_wb.sheetnames[0])
    if sheet_name not in preview_wb.sheetnames:
        raise AggregationError(f"找不到工作表“{sheet_name}”。")
    header_row = prompt_header_row()
    headers, _ = get_headers(preview_wb[sheet_name], header_row)
    print(f"可用列：{'、'.join(headers)}")

    group_raw = prompt_text("请输入分组列（多个列用逗号分隔，例如：项目,规格）")
    sum_raw = prompt_text(
        "请输入需要累加的数值列（多个列用逗号分隔；若仅合并文本行可直接回车）",
        default="",
        required=False,
    )
    policy_raw = prompt_text(
        "其他字段不一致时如何处理：first=保留首个非空值；join=去重拼接",
        default="first",
    ).lower()
    if policy_raw not in {"first", "join"}:
        raise AggregationError("其他字段处理方式只能填写 first 或 join。")

    default_output = str(input_path.with_name(f"{input_path.stem}_汇总.xlsx"))
    output_path = clean_path_text(prompt_text("请输入输出文件完整路径", default_output))
    return {
        "input_path": input_path,
        "output_path": output_path,
        "sheet_name": sheet_name,
        "header_row": header_row,
        "group_names": split_column_names(group_raw),
        "sum_names": split_column_names(sum_raw),
        "text_policy": policy_raw,
        "overwrite": False,
    }


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="按指定列分组合并 Excel 行，并对数值列求和。未带参数运行时进入中文交互模式。"
    )
    parser.add_argument("--input", dest="input_path", help="输入 .xlsx 文件路径")
    parser.add_argument("--output", dest="output_path", help="输出 .xlsx 文件路径")
    parser.add_argument("--sheet", help="需要汇总的工作表名称")
    parser.add_argument("--header-row", type=int, default=1, help="表头所在行号，默认 1")
    parser.add_argument("--group-by", help="分组列名，多个列用逗号分隔")
    parser.add_argument("--sum-cols", default="", help="累加列名，多个列用逗号分隔；可留空")
    parser.add_argument(
        "--text-policy",
        choices=["first", "join"],
        default="first",
        help="其他字段不一致时：first=保留首个非空值；join=去重拼接",
    )
    parser.add_argument("--overwrite", action="store_true", help="允许覆盖已存在的输出文件")
    return parser


def command_line_config(args: argparse.Namespace, parser: argparse.ArgumentParser) -> Dict[str, Any] | None:
    values = [args.input_path, args.output_path, args.sheet, args.group_by]
    if not any(values):
        return None
    missing = []
    if not args.input_path:
        missing.append("--input")
    if not args.output_path:
        missing.append("--output")
    if not args.sheet:
        missing.append("--sheet")
    if not args.group_by:
        missing.append("--group-by")
    if missing:
        parser.error("命令行模式缺少参数：" + "、".join(missing))
    return {
        "input_path": clean_path_text(args.input_path),
        "output_path": clean_path_text(args.output_path),
        "sheet_name": args.sheet,
        "header_row": args.header_row,
        "group_names": split_column_names(args.group_by),
        "sum_names": split_column_names(args.sum_cols),
        "text_policy": args.text_policy,
        "overwrite": args.overwrite,
    }


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    config = command_line_config(args, parser)
    interactive = config is None
    if interactive:
        config = interactive_config()

    result = run_aggregation(**config)
    print("\n处理完成。")
    print(f"原始非空数据行：{result['original_rows']} 行")
    print(f"汇总后数据行：{result['result_rows']} 行")
    print(f"字段冲突：{result['conflicts']} 项（请查看“{result['conflict_sheet']}”工作表）")
    print(f"输出文件：{config['output_path']}")
    if interactive:
        pause_before_exit()
    return 0


def pause_before_exit() -> None:
    """双击 EXE 时保留窗口；命令行或无标准输入环境下不报错。"""
    try:
        input("\n按回车键退出。")
    except EOFError:
        pass


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except AggregationError as exc:
        print(f"\n未完成处理：{exc}")
        if len(sys.argv) == 1:
            pause_before_exit()
        raise SystemExit(2)
    except Exception as exc:
        print("\n未完成处理：发生未预期错误。")
        print(f"原因：{exc}")
        print("请确认输入文件未损坏、输出文件未被 Excel 占用，并检查字段设置后重试。")
        if len(sys.argv) == 1:
            pause_before_exit()
        raise SystemExit(1)
